"""Payroll run lifecycle: draft → submit → approve → export → upload → finalize."""

import calendar
import io
import os
import uuid
from datetime import date
from decimal import Decimal, ROUND_HALF_UP

from fpdf import FPDF
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import (
    Employee,
    EmployeeBankAccount,
    EmployeeCompensation,
    LeaveRequest,
    LeaveType,
    PayrollRun,
    PayrollRunLine,
    Payslip,
    Tenant,
    TenantSetting,
)
from app.services.finance_service import FinanceService, build_salary_breakdown

PAYROLL_SETTINGS_KEY = "payroll_settings"
DEFAULT_PAYROLL_SETTINGS = {
    "working_days_per_month": 30,
    "hr_roles": ["HR Admin", "HR Executive"],
    "finance_roles": ["Finance Admin"],
    "admin_roles": ["Org Admin"],
}

MONTH_NAMES = [
    "", "January", "February", "March", "April", "May", "June",
    "July", "August", "September", "October", "November", "December",
]


def _quantize(value: Decimal) -> Decimal:
    return value.quantize(Decimal("0.01"), ROUND_HALF_UP)


def _sum_dict_values(d: dict) -> Decimal:
    return sum((Decimal(str(v)) for v in d.values()), Decimal("0"))


def recalculate_line(line: PayrollRunLine, working_days: int = 30) -> None:
    """Recompute gross, deductions, and net for a payroll line."""
    days_in_month = Decimal(str(working_days))
    lop_days = Decimal(str(line.lop_days or 0))
    base_gross = Decimal(str(line.base_gross or 0))
    bonus = Decimal(str(line.bonus or 0))

    lop_amount = _quantize((base_gross / days_in_month) * lop_days) if lop_days > 0 else Decimal("0")
    line.lop_amount = lop_amount

    base_deduction_total = _sum_dict_values(line.base_deductions or {})
    # Pro-rate statutory deductions for LOP days
    if lop_days > 0 and days_in_month > 0:
        deduction_factor = (days_in_month - lop_days) / days_in_month
        adjusted_deductions = _quantize(base_deduction_total * deduction_factor)
    else:
        adjusted_deductions = base_deduction_total

    other_earnings = _sum_dict_values(line.other_earnings or {})
    other_deductions = _sum_dict_values(line.other_deductions or {})

    gross_pay = _quantize(base_gross - lop_amount + bonus + other_earnings)
    total_deductions = _quantize(adjusted_deductions + other_deductions)
    net_pay = _quantize(gross_pay - total_deductions)

    line.gross_pay = gross_pay
    line.total_deductions = total_deductions
    line.net_pay = max(net_pay, Decimal("0"))


class PayrollService:
    def __init__(self) -> None:
        self.finance = FinanceService()

    async def get_settings(self, db: AsyncSession, tenant_id: uuid.UUID) -> dict:
        result = await db.execute(
            select(TenantSetting).where(
                TenantSetting.tenant_id == tenant_id,
                TenantSetting.key == PAYROLL_SETTINGS_KEY,
                TenantSetting.is_deleted.is_(False),
            )
        )
        setting = result.scalar_one_or_none()
        if setting and setting.value:
            return {**DEFAULT_PAYROLL_SETTINGS, **setting.value}
        return dict(DEFAULT_PAYROLL_SETTINGS)

    async def upsert_settings(
        self, db: AsyncSession, tenant_id: uuid.UUID, value: dict, user_id: uuid.UUID | None
    ) -> dict:
        merged = {**DEFAULT_PAYROLL_SETTINGS, **value}
        result = await db.execute(
            select(TenantSetting).where(
                TenantSetting.tenant_id == tenant_id,
                TenantSetting.key == PAYROLL_SETTINGS_KEY,
                TenantSetting.is_deleted.is_(False),
            )
        )
        setting = result.scalar_one_or_none()
        if setting:
            setting.value = merged
        else:
            db.add(
                TenantSetting(
                    tenant_id=tenant_id,
                    key=PAYROLL_SETTINGS_KEY,
                    value=merged,
                    created_by=user_id,
                )
            )
        await db.flush()
        return merged

    async def get_run(
        self, db: AsyncSession, tenant_id: uuid.UUID, run_id: uuid.UUID
    ) -> PayrollRun | None:
        result = await db.execute(
            select(PayrollRun).where(
                PayrollRun.id == run_id,
                PayrollRun.tenant_id == tenant_id,
                PayrollRun.is_deleted.is_(False),
            )
        )
        return result.scalar_one_or_none()

    async def list_runs(
        self, db: AsyncSession, tenant_id: uuid.UUID, status: str | None = None
    ) -> list[PayrollRun]:
        q = select(PayrollRun).where(
            PayrollRun.tenant_id == tenant_id, PayrollRun.is_deleted.is_(False)
        )
        if status:
            q = q.where(PayrollRun.status == status)
        result = await db.execute(q.order_by(PayrollRun.year.desc(), PayrollRun.month.desc()))
        return list(result.scalars().all())

    async def _get_lines(self, db: AsyncSession, run_id: uuid.UUID) -> list[PayrollRunLine]:
        result = await db.execute(
            select(PayrollRunLine)
            .where(PayrollRunLine.payroll_run_id == run_id, PayrollRunLine.is_deleted.is_(False))
            .order_by(PayrollRunLine.employee_code)
        )
        return list(result.scalars().all())

    async def _compute_lop_for_period(
        self,
        db: AsyncSession,
        tenant_id: uuid.UUID,
        employee_id: uuid.UUID,
        month: int,
        year: int,
    ) -> tuple[Decimal, dict]:
        """Sum unpaid leave days overlapping the pay period."""
        period_start = date(year, month, 1)
        last_day = calendar.monthrange(year, month)[1]
        period_end = date(year, month, last_day)

        result = await db.execute(
            select(LeaveRequest, LeaveType)
            .join(LeaveType, LeaveRequest.leave_type_id == LeaveType.id)
            .where(
                LeaveRequest.tenant_id == tenant_id,
                LeaveRequest.employee_id == employee_id,
                LeaveRequest.status == "approved",
                LeaveRequest.is_deleted.is_(False),
                LeaveRequest.start_date <= period_end,
                LeaveRequest.end_date >= period_start,
            )
        )
        rows = result.all()
        total_lop = Decimal("0")
        summary: dict = {}
        for req, leave_type in rows:
            if leave_type.is_paid:
                continue
            overlap_start = max(req.start_date, period_start)
            overlap_end = min(req.end_date, period_end)
            if overlap_start > overlap_end:
                continue
            if req.is_half_day and overlap_start == overlap_end:
                days = Decimal("0.5")
            else:
                days = Decimal(str((overlap_end - overlap_start).days + 1))
            total_lop += days
            key = leave_type.name
            summary[key] = float(summary.get(key, 0)) + float(days)
        return total_lop, summary

    async def _bank_snapshot(
        self, db: AsyncSession, tenant_id: uuid.UUID, employee: Employee
    ) -> dict:
        result = await db.execute(
            select(EmployeeBankAccount).where(
                EmployeeBankAccount.tenant_id == tenant_id,
                EmployeeBankAccount.employee_id == employee.id,
                EmployeeBankAccount.is_deleted.is_(False),
            )
        )
        bank = result.scalar_one_or_none()
        if bank:
            return {
                "account_holder_name": bank.account_holder_name,
                "account_number": bank.account_number,
                "ifsc_code": bank.ifsc_code,
                "bank_name": bank.bank_name,
                "branch_name": bank.branch_name,
            }
        return {
            "account_holder_name": f"{employee.first_name} {employee.last_name}",
            "account_number": "",
            "ifsc_code": "",
            "bank_name": "",
            "branch_name": "",
        }

    def _recalc_run_totals(self, run: PayrollRun, lines: list[PayrollRunLine]) -> None:
        run.total_gross = sum((line.gross_pay for line in lines), Decimal("0"))
        run.total_net = sum((line.net_pay for line in lines), Decimal("0"))
        run.employee_count = len(lines)

    async def create_draft(
        self,
        db: AsyncSession,
        *,
        tenant_id: uuid.UUID,
        month: int,
        year: int,
        user_id: uuid.UUID | None,
        notes: str | None = None,
    ) -> PayrollRun:
        existing = await db.execute(
            select(PayrollRun).where(
                PayrollRun.tenant_id == tenant_id,
                PayrollRun.month == month,
                PayrollRun.year == year,
                PayrollRun.is_deleted.is_(False),
            )
        )
        if existing.scalar_one_or_none():
            raise ValueError(f"Payroll run for {MONTH_NAMES[month]} {year} already exists")

        settings = await self.get_settings(db, tenant_id)
        working_days = int(settings.get("working_days_per_month", 30))

        run = PayrollRun(
            tenant_id=tenant_id,
            month=month,
            year=year,
            status="draft",
            notes=notes,
            created_by=user_id,
        )
        db.add(run)
        await db.flush()

        emp_result = await db.execute(
            select(Employee).where(
                Employee.tenant_id == tenant_id,
                Employee.is_deleted.is_(False),
                Employee.employment_status == "active",
            ).order_by(Employee.employee_code)
        )
        employees = list(emp_result.scalars().all())
        lines: list[PayrollRunLine] = []

        for emp in employees:
            comp = await self.finance.get_compensation(db, tenant_id, emp.id)
            base_gross = Decimal("0")
            base_deductions: dict = {}
            if comp:
                enriched = self.finance.enrichment_for_compensation(comp)
                base_gross = Decimal(str(enriched["gross_monthly"] or 0))
                base_deductions = {
                    k: v for k, v in (enriched["deductions"] or {}).items() if k != "tax_regime"
                }
            elif comp and comp.ctc_annual:
                _, deductions, gross, _, _ = build_salary_breakdown(comp.ctc_annual)
                base_gross = gross or Decimal("0")
                base_deductions = deductions

            lop_days, leave_summary = await self._compute_lop_for_period(
                db, tenant_id, emp.id, month, year
            )
            bank_snapshot = await self._bank_snapshot(db, tenant_id, emp)

            line = PayrollRunLine(
                tenant_id=tenant_id,
                payroll_run_id=run.id,
                employee_id=emp.id,
                employee_code=emp.employee_code,
                employee_name=f"{emp.first_name} {emp.last_name}",
                base_gross=base_gross,
                base_deductions=base_deductions,
                lop_days=lop_days,
                bonus=Decimal("0"),
                leave_summary=leave_summary,
                bank_snapshot=bank_snapshot,
                created_by=user_id,
            )
            recalculate_line(line, working_days)
            lines.append(line)
            db.add(line)

        self._recalc_run_totals(run, lines)
        await db.flush()
        return run

    async def update_line(
        self,
        db: AsyncSession,
        *,
        tenant_id: uuid.UUID,
        run_id: uuid.UUID,
        line_id: uuid.UUID,
        updates: dict,
    ) -> PayrollRunLine:
        run = await self.get_run(db, tenant_id, run_id)
        if not run:
            raise ValueError("Payroll run not found")
        if run.status not in ("draft", "rejected"):
            raise ValueError("Can only edit draft or rejected payroll runs")

        result = await db.execute(
            select(PayrollRunLine).where(
                PayrollRunLine.id == line_id,
                PayrollRunLine.payroll_run_id == run_id,
                PayrollRunLine.tenant_id == tenant_id,
            )
        )
        line = result.scalar_one_or_none()
        if not line:
            raise ValueError("Payroll line not found")

        for field in ("lop_days", "bonus", "other_earnings", "other_deductions", "notes"):
            if field in updates and updates[field] is not None:
                setattr(line, field, updates[field])

        settings = await self.get_settings(db, tenant_id)
        recalculate_line(line, int(settings.get("working_days_per_month", 30)))

        lines = await self._get_lines(db, run_id)
        self._recalc_run_totals(run, lines)
        await db.flush()
        return line

    async def submit_for_review(
        self, db: AsyncSession, *, tenant_id: uuid.UUID, run_id: uuid.UUID, user_id: uuid.UUID
    ) -> PayrollRun:
        run = await self.get_run(db, tenant_id, run_id)
        if not run:
            raise ValueError("Payroll run not found")
        if run.status not in ("draft", "rejected"):
            raise ValueError("Only draft or rejected runs can be submitted")
        run.status = "submitted"
        run.submitted_by = user_id
        run.submitted_at = date.today()
        run.finance_notes = None
        await db.flush()
        return run

    async def approve(
        self, db: AsyncSession, *, tenant_id: uuid.UUID, run_id: uuid.UUID, user_id: uuid.UUID
    ) -> PayrollRun:
        run = await self.get_run(db, tenant_id, run_id)
        if not run:
            raise ValueError("Payroll run not found")
        if run.status != "submitted":
            raise ValueError("Only submitted runs can be approved")
        run.status = "approved"
        run.approved_by = user_id
        run.approved_at = date.today()
        await db.flush()
        return run

    async def reject(
        self,
        db: AsyncSession,
        *,
        tenant_id: uuid.UUID,
        run_id: uuid.UUID,
        user_id: uuid.UUID,
        finance_notes: str,
    ) -> PayrollRun:
        run = await self.get_run(db, tenant_id, run_id)
        if not run:
            raise ValueError("Payroll run not found")
        if run.status != "submitted":
            raise ValueError("Only submitted runs can be rejected")
        run.status = "rejected"
        run.finance_notes = finance_notes
        run.approved_by = user_id
        run.approved_at = date.today()
        await db.flush()
        return run

    async def export_bank_excel(
        self, db: AsyncSession, tenant_id: uuid.UUID, run_id: uuid.UUID
    ) -> tuple[bytes, str]:
        run = await self.get_run(db, tenant_id, run_id)
        if not run:
            raise ValueError("Payroll run not found")
        if run.status not in ("approved", "processing", "completed"):
            raise ValueError("Payroll must be approved before export")

        lines = await self._get_lines(db, run_id)
        wb = Workbook()
        ws = wb.active
        ws.title = "Payroll Batch"
        headers = [
            "Employee Code",
            "Employee Name",
            "Account Holder",
            "Account Number",
            "IFSC",
            "Bank Name",
            "Branch",
            "Net Pay",
            "Payment Status",
            "Payment Reference",
        ]
        ws.append(headers)
        for line in lines:
            bank = line.bank_snapshot or {}
            ws.append([
                line.employee_code,
                line.employee_name,
                bank.get("account_holder_name", ""),
                bank.get("account_number", ""),
                bank.get("ifsc_code", ""),
                bank.get("bank_name", ""),
                bank.get("branch_name", ""),
                float(line.net_pay),
                line.payment_status or "",
                line.payment_reference or "",
            ])

        if run.status == "approved":
            run.status = "processing"
            await db.flush()

        buf = io.BytesIO()
        wb.save(buf)
        filename = f"payroll_{run.year}_{run.month:02d}_batch.xlsx"
        return buf.getvalue(), filename

    async def import_payment_status(
        self, db: AsyncSession, tenant_id: uuid.UUID, run_id: uuid.UUID, file_bytes: bytes
    ) -> int:
        run = await self.get_run(db, tenant_id, run_id)
        if not run:
            raise ValueError("Payroll run not found")
        if run.status not in ("approved", "processing"):
            raise ValueError("Payroll must be approved or processing to upload payment status")

        wb = load_workbook(io.BytesIO(file_bytes), read_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 2:
            raise ValueError("Excel file has no data rows")

        header = [str(h).strip().lower() if h else "" for h in rows[0]]
        code_idx = header.index("employee code") if "employee code" in header else 0
        status_idx = header.index("payment status") if "payment status" in header else -1
        ref_idx = header.index("payment reference") if "payment reference" in header else -1

        if status_idx < 0:
            raise ValueError("Excel must contain a 'Payment Status' column")

        lines = await self._get_lines(db, run_id)
        line_by_code = {line.employee_code: line for line in lines}
        updated = 0

        for row in rows[1:]:
            if not row or not row[code_idx]:
                continue
            code = str(row[code_idx]).strip()
            line = line_by_code.get(code)
            if not line:
                continue
            status_val = row[status_idx] if status_idx < len(row) else None
            if status_val:
                line.payment_status = str(status_val).strip()
                updated += 1
            if ref_idx >= 0 and ref_idx < len(row) and row[ref_idx]:
                line.payment_reference = str(row[ref_idx]).strip()

        await db.flush()
        return updated

    def _generate_payslip_pdf(
        self,
        *,
        company_name: str,
        employee_name: str,
        employee_code: str,
        month: int,
        year: int,
        line: PayrollRunLine,
    ) -> bytes:
        pdf = FPDF()
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.add_page()
        pdf.set_font("Helvetica", "B", 16)
        pdf.cell(0, 10, company_name, ln=True, align="C")
        pdf.set_font("Helvetica", "", 12)
        pdf.cell(0, 8, f"Payslip - {MONTH_NAMES[month]} {year}", ln=True, align="C")
        pdf.ln(5)
        pdf.set_font("Helvetica", "", 10)
        pdf.cell(0, 6, f"Employee: {employee_name} ({employee_code})", ln=True)
        pdf.ln(3)

        pdf.set_font("Helvetica", "B", 11)
        pdf.cell(0, 7, "Earnings", ln=True)
        pdf.set_font("Helvetica", "", 10)
        pdf.cell(100, 6, "Base Gross")
        pdf.cell(0, 6, f"{float(line.base_gross):,.2f}", ln=True)
        if float(line.bonus) > 0:
            pdf.cell(100, 6, "Bonus")
            pdf.cell(0, 6, f"{float(line.bonus):,.2f}", ln=True)
        for key, val in (line.other_earnings or {}).items():
            pdf.cell(100, 6, key.replace("_", " ").title())
            pdf.cell(0, 6, f"{float(val):,.2f}", ln=True)
        if float(line.lop_amount) > 0:
            pdf.cell(100, 6, f"LOP ({line.lop_days} days)")
            pdf.cell(0, 6, f"-{float(line.lop_amount):,.2f}", ln=True)
        pdf.cell(100, 6, "Gross Pay")
        pdf.cell(0, 6, f"{float(line.gross_pay):,.2f}", ln=True)

        pdf.ln(3)
        pdf.set_font("Helvetica", "B", 11)
        pdf.cell(0, 7, "Deductions", ln=True)
        pdf.set_font("Helvetica", "", 10)
        for key, val in (line.base_deductions or {}).items():
            pdf.cell(100, 6, key.replace("_", " ").title())
            pdf.cell(0, 6, f"{float(val):,.2f}", ln=True)
        for key, val in (line.other_deductions or {}).items():
            pdf.cell(100, 6, key.replace("_", " ").title())
            pdf.cell(0, 6, f"{float(val):,.2f}", ln=True)

        pdf.ln(3)
        pdf.set_font("Helvetica", "B", 12)
        pdf.cell(100, 8, "Net Pay")
        pdf.cell(0, 8, f"INR {float(line.net_pay):,.2f}", ln=True)

        if line.payment_status:
            pdf.ln(3)
            pdf.set_font("Helvetica", "", 10)
            pdf.cell(0, 6, f"Payment Status: {line.payment_status}", ln=True)
            if line.payment_reference:
                pdf.cell(0, 6, f"Reference: {line.payment_reference}", ln=True)

        return bytes(pdf.output())

    async def finalize(
        self, db: AsyncSession, *, tenant_id: uuid.UUID, run_id: uuid.UUID, user_id: uuid.UUID | None
    ) -> PayrollRun:
        run = await self.get_run(db, tenant_id, run_id)
        if not run:
            raise ValueError("Payroll run not found")
        if run.status not in ("approved", "processing"):
            raise ValueError("Payroll must be approved or processing to finalize")

        tenant = await db.get(Tenant, tenant_id)
        company_name = tenant.name if tenant else "Company"

        lines = await self._get_lines(db, run_id)
        upload_dir = os.path.join("uploads", str(tenant_id), "payslips")
        os.makedirs(upload_dir, exist_ok=True)

        for line in lines:
            existing = await db.execute(
                select(Payslip).where(
                    Payslip.tenant_id == tenant_id,
                    Payslip.employee_id == line.employee_id,
                    Payslip.month == run.month,
                    Payslip.year == run.year,
                )
            )
            if existing.scalar_one_or_none():
                continue

            earnings = {
                "base_gross": float(line.base_gross),
                "bonus": float(line.bonus),
                **{k: float(v) for k, v in (line.other_earnings or {}).items()},
            }
            if float(line.lop_amount) > 0:
                earnings["lop_adjustment"] = -float(line.lop_amount)

            deductions = {
                **{k: float(v) for k, v in (line.base_deductions or {}).items()},
                **{k: float(v) for k, v in (line.other_deductions or {}).items()},
            }

            pdf_bytes = self._generate_payslip_pdf(
                company_name=company_name,
                employee_name=line.employee_name,
                employee_code=line.employee_code,
                month=run.month,
                year=run.year,
                line=line,
            )
            filename = f"payslip_{line.employee_code}_{run.year}_{run.month:02d}.pdf"
            filepath = os.path.join(upload_dir, filename)
            with open(filepath, "wb") as f:
                f.write(pdf_bytes)

            payslip = Payslip(
                tenant_id=tenant_id,
                employee_id=line.employee_id,
                month=run.month,
                year=run.year,
                gross_pay=line.gross_pay,
                net_pay=line.net_pay,
                earnings=earnings,
                deductions=deductions,
                file_url=filepath,
                created_by=user_id,
            )
            db.add(payslip)

        run.status = "completed"
        await db.flush()
        return run

    async def upsert_bank_account(
        self,
        db: AsyncSession,
        *,
        tenant_id: uuid.UUID,
        employee_id: uuid.UUID,
        user_id: uuid.UUID | None,
        account_holder_name: str,
        account_number: str,
        ifsc_code: str,
        bank_name: str,
        branch_name: str | None = None,
    ) -> EmployeeBankAccount:
        result = await db.execute(
            select(EmployeeBankAccount).where(
                EmployeeBankAccount.tenant_id == tenant_id,
                EmployeeBankAccount.employee_id == employee_id,
                EmployeeBankAccount.is_deleted.is_(False),
            )
        )
        existing = result.scalar_one_or_none()
        if existing:
            existing.account_holder_name = account_holder_name
            existing.account_number = account_number
            existing.ifsc_code = ifsc_code.upper()
            existing.bank_name = bank_name
            existing.branch_name = branch_name
            await db.flush()
            return existing

        bank = EmployeeBankAccount(
            tenant_id=tenant_id,
            employee_id=employee_id,
            account_holder_name=account_holder_name,
            account_number=account_number,
            ifsc_code=ifsc_code.upper(),
            bank_name=bank_name,
            branch_name=branch_name,
            created_by=user_id,
        )
        db.add(bank)
        await db.flush()
        return bank

    async def get_bank_account(
        self, db: AsyncSession, tenant_id: uuid.UUID, employee_id: uuid.UUID
    ) -> EmployeeBankAccount | None:
        result = await db.execute(
            select(EmployeeBankAccount).where(
                EmployeeBankAccount.tenant_id == tenant_id,
                EmployeeBankAccount.employee_id == employee_id,
                EmployeeBankAccount.is_deleted.is_(False),
            )
        )
        return result.scalar_one_or_none()

    async def refresh_draft_lines(
        self, db: AsyncSession, *, tenant_id: uuid.UUID, run_id: uuid.UUID
    ) -> PayrollRun:
        """Re-sync all employees and leave data into an existing draft."""
        run = await self.get_run(db, tenant_id, run_id)
        if not run:
            raise ValueError("Payroll run not found")
        if run.status not in ("draft", "rejected"):
            raise ValueError("Can only refresh draft or rejected runs")

        settings = await self.get_settings(db, tenant_id)
        working_days = int(settings.get("working_days_per_month", 30))

        existing_lines = {line.employee_id: line for line in await self._get_lines(db, run_id)}

        emp_result = await db.execute(
            select(Employee).where(
                Employee.tenant_id == tenant_id,
                Employee.is_deleted.is_(False),
                Employee.employment_status == "active",
            )
        )
        active_ids = set()
        for emp in emp_result.scalars().all():
            active_ids.add(emp.id)
            line = existing_lines.get(emp.id)
            if not line:
                comp = await self.finance.get_compensation(db, tenant_id, emp.id)
                base_gross = Decimal("0")
                base_deductions: dict = {}
                if comp:
                    enriched = self.finance.enrichment_for_compensation(comp)
                    base_gross = Decimal(str(enriched["gross_monthly"] or 0))
                    base_deductions = {
                        k: v for k, v in (enriched["deductions"] or {}).items() if k != "tax_regime"
                    }
                lop_days, leave_summary = await self._compute_lop_for_period(
                    db, tenant_id, emp.id, run.month, run.year
                )
                line = PayrollRunLine(
                    tenant_id=tenant_id,
                    payroll_run_id=run.id,
                    employee_id=emp.id,
                    employee_code=emp.employee_code,
                    employee_name=f"{emp.first_name} {emp.last_name}",
                    base_gross=base_gross,
                    base_deductions=base_deductions,
                    lop_days=lop_days,
                    leave_summary=leave_summary,
                    bank_snapshot=await self._bank_snapshot(db, tenant_id, emp),
                    created_by=run.created_by,
                )
                db.add(line)
                existing_lines[emp.id] = line
            else:
                lop_days, leave_summary = await self._compute_lop_for_period(
                    db, tenant_id, emp.id, run.month, run.year
                )
                line.lop_days = lop_days
                line.leave_summary = leave_summary
                line.bank_snapshot = await self._bank_snapshot(db, tenant_id, emp)
                line.employee_code = emp.employee_code
                line.employee_name = f"{emp.first_name} {emp.last_name}"

            recalculate_line(line, working_days)

        for emp_id, line in existing_lines.items():
            if emp_id not in active_ids:
                line.is_deleted = True

        lines = [l for l in existing_lines.values() if not l.is_deleted]
        self._recalc_run_totals(run, lines)
        await db.flush()
        return run
